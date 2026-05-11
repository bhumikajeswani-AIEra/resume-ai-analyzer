import io
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SECTION_HEADERS = {
    "experience", "work experience", "professional experience", "employment",
    "education", "academic background",
    "skills", "technical skills", "core skills", "key skills",
    "certifications", "certificates", "licenses",
    "projects", "key projects", "personal projects",
    "summary", "professional summary", "objective", "profile", "about",
    "achievements", "accomplishments", "awards",
    "publications", "research", "interests", "languages", "references",
    "internships", "volunteer", "extracurricular",
}

TEMPLATE_STYLES = {
    "technical": {"heading_color": RGBColor(0x1e, 0x40, 0xaf), "name_size": 16, "border": False, "caps": True},
    "executive": {"heading_color": RGBColor(0x1f, 0x2d, 0x3d), "name_size": 18, "border": True,  "caps": True},
    "creative":  {"heading_color": RGBColor(0x03, 0x69, 0xa1), "name_size": 18, "border": False, "caps": False},
    "academic":  {"heading_color": RGBColor(0x33, 0x33, 0x33), "name_size": 16, "border": False, "caps": True},
    "general":   {"heading_color": RGBColor(0x22, 0x22, 0x22), "name_size": 16, "border": True,  "caps": True},
}


def _is_section_header(line: str) -> bool:
    stripped = line.strip().rstrip(":").lower()
    if stripped in SECTION_HEADERS:
        return True
    if line.strip().isupper() and 2 < len(line.strip()) < 50:
        return True
    return False


def _is_contact_line(line: str) -> bool:
    l = line.lower()
    return bool(
        "@" in l or
        re.search(r"\+?\d[\d\s\-]{7,}", line) or
        "linkedin" in l or "github" in l or
        l.strip().startswith(("http", "www"))
    )


def _set_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_docx(resume_text: str, template_name: str) -> bytes:
    style = TEMPLATE_STYLES.get(template_name, TEMPLATE_STYLES["general"])
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    lines = resume_text.split("\n")
    first_nonempty = True

    for line in lines:
        stripped = line.strip()

        if not stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            continue

        # First non-empty line = name
        if first_nonempty:
            first_nonempty = False
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.bold = True
            run.font.size = Pt(style["name_size"])
            run.font.color.rgb = style["heading_color"]
            p.paragraph_format.space_after = Pt(2)
            continue

        # Contact / header info lines
        if _is_contact_line(stripped):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            p.paragraph_format.space_after = Pt(1)
            continue

        # Section header
        if _is_section_header(stripped):
            p = doc.add_paragraph()
            label = stripped.upper() if style["caps"] else stripped
            run = p.add_run(label)
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = style["heading_color"]
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            if style["border"]:
                _set_bottom_border(p)
            continue

        # Bullet point
        if stripped.startswith(("•", "-", "–", "▪", "*", "○")):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(stripped.lstrip("•-–▪*○ ").strip())
            run.font.size = Pt(10)
            p.paragraph_format.space_after = Pt(1)
            continue

        # Regular text
        p = doc.add_paragraph()
        run = p.add_run(stripped)
        run.font.size = Pt(10)
        p.paragraph_format.space_after = Pt(2)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
